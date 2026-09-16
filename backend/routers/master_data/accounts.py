"""Modular routes: accounts"""
from router_dependencies import (
    ACCESS_TOKEN_EXPIRE_HOURS,
    ADMIN_LEVEL,
    APIRouter,
    API_PREFIX,
    APP_TITLE,
    Account,
    AccountCreate,
    DatabaseClient,
    CHART_OF_ACCOUNTS,
    EXAMPLE_ROWS_BUMDES,
    EXAMPLE_ROWS_UNIT,
    COOKIE_NAME,
    CORSMiddleware,
    ChangePasswordRequest,
    Depends,
    FastAPI,
    File,
    HTMLResponse,
    HTTPException,
    List,
    Mitra,
    MitraCreate,
    Optional,
    P,
    Paragraph,
    PasswordResetRequest,
    Path,
    ProfileUpdateRequest,
    Query,
    READONLY_ROLES,
    READ_LEVEL,
    REPORT_READ_LEVEL,
    ROOT_DIR,
    Request,
    Response,
    RevenueShare,
    RevenueShareCreate,
    Spacer,
    StreamingResponse,
    Table,
    TableStyle,
    Transaction,
    TransactionCreate,
    UNIT_USAHA_SEED,
    VALID_CATEGORIES,
    UnitUsaha,
    UnitUsahaCreate,
    UploadFile,
    User,
    UserCreate,
    UserLogin,
    UserOut,
    UserRole,
    WRITE_LEVEL,
    _CELL_BOLD,
    _CELL_BOLD_RIGHT,
    _CELL_RIGHT,
    _CELL_STYLE,
    _STYLES,
    _arus_kas,
    _calc_balances,
    _calc_balances_before,
    _get_accounts_map,
    _group_from_unit,
    _laba_rugi,
    _ledger_data,
    _neraca,
    _pdf_response,
    _per_unit_report,
    _perubahan_ekuitas,
    _section_row,
    _sig_flow,
    _table_style,
    app,
    client,
    close_database,
    cm,
    colors,
    create_access_token,
    datetime,
    db,
    fmt_rp,
    get_current_user_payload,
    hash_password,
    io,
    load_dotenv,
    logging,
    now_utc,
    os,
    pdf_header,
    require_not_readonly,
    require_password_ready,
    require_roles,
    router,
    scope_unit_for_pengelola,
    seed_database,
    timezone,
    user_from_payload,
    uuid,
    verify_password
)

router = APIRouter(prefix=API_PREFIX)

@router.get("/accounts/template")
async def download_accounts_template(_: dict = Depends(require_roles(*ADMIN_LEVEL))):
    """Excel template dengan sheet Instruksi + BUMDES + UU01..UU06."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    import io as _io

    wb = Workbook()
    # ------ Sheet 1: Instruksi ------
    ws = wb.active
    ws.title = "Instruksi"
    header_font = Font(bold=True, size=14, color="FFFFFF")
    section_font = Font(bold=True, size=11, color="2E4F32")
    fill_head = PatternFill("solid", fgColor="2E4F32")

    ws["A1"] = "Template Import Kode Akun (COA) - BUMDES Karya Raharja"
    ws["A1"].font = header_font
    ws["A1"].fill = fill_head
    ws.merge_cells("A1:E1")
    ws.row_dimensions[1].height = 24

    lines = [
        "",
        "CARA PENGISIAN:",
        "1. Isi baris di sheet 'BUMDES' untuk akun tingkat BUMDES; sheet 'UU01' - 'UU06' untuk tiap unit usaha.",
        "2. Jangan mengubah judul sheet atau nama kolom (baris header).",
        "3. Kolom 'code' dapat berupa format apa saja (mis. 1.1.01.01, 1.1.01.01, dst) - unik dalam kelompok.",
        "4. Sistem tidak lagi hardcoded kode akun. Fungsi akun ditentukan dari 'category' & 'subcategory'.",
        "",
        "KOLOM WAJIB:",
        "  code            - kode unik dalam grup (contoh: 1.1.01.01 atau 1.1.01.01)",
        "  name            - nama akun (contoh: Kas Bendahara)",
        "  category        - kategori akun (lihat daftar di bawah)",
        "  subcategory     - sub kategori (lihat daftar di bawah - HARUS sesuai dengan category)",
        "  normal_balance  - 'debit' atau 'kredit'",
        "",
        "DAFTAR KATEGORI & SUB-KATEGORI VALID:",
    ]
    row = 2
    for line in lines:
        ws.cell(row=row, column=1, value=line)
        if line.endswith(":"):
            ws.cell(row=row, column=1).font = section_font
        row += 1

    for cat, subs in VALID_CATEGORIES.items():
        ws.cell(row=row, column=1, value=f"  {cat}")
        ws.cell(row=row, column=1).font = Font(bold=True, color="8CA650")
        ws.cell(row=row, column=2, value=", ".join(subs))
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="ATURAN SALDO NORMAL:")
    ws.cell(row=row, column=1).font = section_font
    row += 1
    for line in [
        "  aset           -> debit",
        "  kewajiban      -> kredit",
        "  ekuitas        -> kredit",
        "  pendapatan     -> kredit",
        "  hpp            -> debit",
        "  beban          -> debit",
    ]:
        ws.cell(row=row, column=1, value=line)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="CATATAN PENTING:")
    ws.cell(row=row, column=1).font = section_font
    row += 1
    notes = [
        "* Akun Kas & Bank WAJIB pakai subcategory 'kas_bank' agar dikenali di Laporan Arus Kas.",
        "* Akun Laba Ditahan pakai subcategory 'saldo_laba' agar terhitung di Perubahan Ekuitas.",
        "* Kode boleh sama antar grup (mis. BUMDES 1.1.01.01 = Kas Bendahara, UU01 1.1.01.01 = Kas UU01).",
        "* Baris kosong akan dilewati saat import.",
        "* Baris dengan kode duplikat dalam grup yang sama akan dilewati (tidak menimpa akun eksisting).",
    ]
    for n in notes:
        ws.cell(row=row, column=1, value=n)
        row += 1

    for col in "ABCDE":
        ws.column_dimensions[col].width = 30

    # ------ Sheets: BUMDES + UU01..UU06 ------
    def add_data_sheet(name: str, examples):
        s = wb.create_sheet(name)
        headers = ["code", "name", "category", "subcategory", "normal_balance"]
        for i, h in enumerate(headers, start=1):
            c = s.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill_head
            c.alignment = Alignment(horizontal="center")
        for r, row_data in enumerate(examples, start=2):
            for i, v in enumerate(row_data, start=1):
                s.cell(row=r, column=i, value=v)
        # column widths
        widths = [16, 40, 16, 30, 16]
        for i, w in enumerate(widths, start=1):
            s.column_dimensions[get_column_letter(i)].width = w
        s.row_dimensions[1].height = 20
        s.freeze_panes = "A2"

    add_data_sheet("BUMDES", EXAMPLE_ROWS_BUMDES)
    for uu in ("UU01", "UU02", "UU03", "UU04", "UU05", "UU06"):
        add_data_sheet(uu, EXAMPLE_ROWS_UNIT[uu])

    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from fastapi.responses import StreamingResponse as _SR
    return _SR(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-Kode-Akun.xlsx"'},
    )

@router.post("/accounts/import")
async def import_accounts(file: UploadFile = File(...), _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    """Import kode akun dari Excel multi-sheet (BUMDES, UU01..UU06)."""
    from openpyxl import load_workbook
    import io as _io

    valid_groups = ("BUMDES", "UU01", "UU02", "UU03", "UU04", "UU05", "UU06")
    content = await file.read()
    try:
        wb = load_workbook(filename=_io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"File tidak valid: {e}")

    inserted = 0
    skipped = 0
    errors: list = []

    for sheet_name in wb.sheetnames:
        if sheet_name not in valid_groups:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        # Detect header
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        expected = ["code", "name", "category", "subcategory", "normal_balance"]
        if header[:5] != expected:
            errors.append(f"Sheet '{sheet_name}': header tidak sesuai (harus: {', '.join(expected)})")
            continue
        for idx, row in enumerate(rows[1:], start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row[:5]):
                continue
            code = (str(row[0]).strip() if row[0] is not None else "")
            name = (str(row[1]).strip() if row[1] is not None else "")
            category = (str(row[2]).strip().lower() if row[2] is not None else "")
            subcategory = (str(row[3]).strip().lower() if row[3] is not None else "")
            normal_balance = (str(row[4]).strip().lower() if row[4] is not None else "")

            if not (code and name and category and normal_balance):
                errors.append(f"{sheet_name} baris {idx}: kolom wajib kosong")
                skipped += 1
                continue
            if category not in VALID_CATEGORIES:
                errors.append(f"{sheet_name} baris {idx}: kategori '{category}' tidak valid")
                skipped += 1
                continue
            if subcategory and subcategory not in VALID_CATEGORIES[category]:
                errors.append(f"{sheet_name} baris {idx}: subkategori '{subcategory}' tidak valid untuk kategori '{category}'")
                skipped += 1
                continue
            if normal_balance not in ("debit", "kredit"):
                errors.append(f"{sheet_name} baris {idx}: normal_balance harus 'debit' atau 'kredit'")
                skipped += 1
                continue
            # Duplicate check
            if await db.accounts.select_one({"code": code, "group": sheet_name}):
                skipped += 1
                continue
            acc = Account(
                code=code, name=name, category=category,
                subcategory=subcategory, normal_balance=normal_balance,
                group=sheet_name,
            )
            await db.accounts.create(acc.model_dump())
            inserted += 1

    return {"inserted": inserted, "skipped": skipped, "errors": errors[:50]}

@router.post("/transaction-types/import")
async def import_transaction_types(file: UploadFile = File(...), _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    from openpyxl import load_workbook
    import uuid
    content = await file.read()
    ws = load_workbook(filename=__import__("io").BytesIO(content), read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows: raise HTTPException(400, "File kosong")
    headers = [str(v or "").strip().lower() for v in rows[0]]
    required = {"code", "name", "debit", "credit"}
    if not required.issubset(headers): raise HTTPException(400, "Kolom wajib: code, name, debit, credit")
    inserted = skipped = 0
    for values in rows[1:]:
        item = dict(zip(headers, values))
        if not item.get("code") or not item.get("name"): continue
        if await db.transaction_types.select_one({"code": str(item["code"])}): skipped += 1; continue
        await db.transaction_types.create({"id": str(uuid.uuid4()), "code": str(item["code"]).strip(), "name": str(item["name"]).strip(), "debit": str(item.get("debit") or "").strip(), "credit": str(item.get("credit") or "").strip(), "group": str(item.get("group") or "BUMDES").strip()})
        inserted += 1
    return {"inserted": inserted, "skipped": skipped, "errors": []}

@router.get("/transaction-types/template")
async def transaction_types_template(group: str = Query("BUMDES"), _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import io as _io
    valid_groups = {"BUMDES", "UU01", "UU02", "UU03", "UU04", "UU05", "UU06"}
    if group not in valid_groups:
        raise HTTPException(400, "Kelompok tidak valid")
    active_accounts = await db.accounts.select({"group": group}, {"_id": 0}).order("code", 1).all(10000)
    if not active_accounts:
        raise HTTPException(400, f"Belum ada kode akun aktif untuk kelompok {group}")
    wb = Workbook()
    instructions = wb.active
    instructions.title = "Instruksi"
    instructions.append(["PANDUAN PENGISIAN JENIS TRANSAKSI"])
    instructions.append(["1. Isi satu jenis transaksi pada setiap baris di sheet Jenis Transaksi."])
    instructions.append(["2. Kolom code wajib unik dan tanpa spasi, contoh: penjualan_kios."])
    instructions.append(["3. Kolom name berisi nama transaksi yang mudah dipahami."])
    instructions.append(["4. Kolom debit dan credit berisi kode akun dari kelompok yang dipilih."])
    instructions.append(["5. Kolom group diisi BUMDES atau kode unit UU01 sampai UU06."])
    instructions.append(["6. Jangan mengubah nama kolom dan hapus baris contoh sebelum import."])
    instructions.append([""])
    instructions.append([f"REFERENSI KODE AKUN AKTIF DAN VALID - KELOMPOK {group}"])
    instructions.append(["Gunakan kode akun pada daftar berikut untuk kolom debit dan credit."])
    instructions.append(["Kode Akun", "Nama Akun"])
    for account in active_accounts:
        instructions.append([account.get("code", ""), account.get("name", "")])
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 48
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="2E4F32")
    for cell in instructions[9]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4F32")
    for cell in instructions[11]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4F32")
    for row in instructions.iter_rows():
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Jenis Transaksi")
    ws.append(["code", "name", "debit", "credit", "group"])
    debit_account = next((a for a in active_accounts if a.get("normal_balance") == "debit"), active_accounts[0])
    credit_account = next((a for a in active_accounts if a.get("normal_balance") == "kredit"), active_accounts[-1])
    ws.append(["contoh_transaksi", "Contoh Transaksi", debit_account.get("code", ""), credit_account.get("code", ""), group])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4F32")
    for column, width in {"A": 24, "B": 32, "C": 18, "D": 18, "E": 14}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    buf = _io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="Template-Jenis-Transaksi.xlsx"'})

@router.get("/master-data/export")
async def export_master_data(group: str = Query("BUMDES"), section: str = Query("all"), _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    """Export kode akun atau jenis transaksi untuk satu kelompok terpilih."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import io as _io

    valid_groups = {"BUMDES", "UU01", "UU02", "UU03", "UU04", "UU05", "UU06"}
    if group not in valid_groups:
        raise HTTPException(400, "Kelompok tidak valid")
    if section not in {"all", "accounts", "transaction-types"}:
        raise HTTPException(400, "Jenis export tidak valid")
    wb = Workbook()
    ws_accounts = wb.active
    ws_accounts.title = "Kode Akun"
    ws_accounts.append(["code", "name", "category", "subcategory", "normal_balance", "group"])
    ws_types = wb.create_sheet("Jenis Transaksi")
    ws_types.append(["code", "name", "debit", "credit", "group"])
    for ws in (ws_accounts, ws_types):
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E4F32")
        ws.freeze_panes = "A2"
    for account in await db.accounts.select({"group": group}, {"_id": 0}).order("code", 1).all(10000):
        ws_accounts.append([account.get(k, "") for k in ("code", "name", "category", "subcategory", "normal_balance", "group")])
    for tx_type in await db.transaction_types.select({"group": group}, {"_id": 0}).order("code", 1).all(10000):
        ws_types.append([tx_type.get(k, "") for k in ("code", "name", "debit", "credit", "group")])
    if section == "accounts":
        del wb["Jenis Transaksi"]
    elif section == "transaction-types":
        del wb["Kode Akun"]
        wb.active = wb["Jenis Transaksi"]
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f'Master-Data-{"Kode-Akun" if section == "accounts" else "Jenis-Transaksi" if section == "transaction-types" else "Lengkap"}-{group}.xlsx'
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})

@router.delete("/accounts/reset-all")
async def reset_all_accounts(confirm: str = "", _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    """Hapus SEMUA kode akun (butuh confirm=YES). Transaksi tidak dihapus."""
    if confirm != "YES":
        raise HTTPException(400, "Parameter confirm=YES wajib untuk operasi ini")
    r = await db.accounts.remove_many({})
    return {"deleted": r.deleted_count}

@router.get("/accounts", response_model=List[Account])
async def list_accounts(_: dict = Depends(get_current_user_payload)):
    docs = await db.accounts.select({}, {"_id": 0}).order("code", 1).all(500)
    return [Account(**d) for d in docs]

@router.post("/accounts", response_model=Account)
async def create_account(payload: AccountCreate, _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    grp = payload.group or "BUMDES"
    if await db.accounts.select_one({"code": payload.code, "group": grp}):
        raise HTTPException(status_code=400, detail=f"Kode akun {payload.code} sudah ada di kelompok {grp}")
    acc = Account(**payload.model_dump())
    await db.accounts.create(acc.model_dump())
    return acc

@router.put("/accounts/{code}", response_model=Account)
async def update_account(code: str, payload: AccountCreate, group: str = "BUMDES",
                         _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    existing = await db.accounts.select_one({"code": code, "group": group}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail=f"Kode akun {code} tidak ditemukan di kelompok {group}")
    data = payload.model_dump()
    target_group = data.get("group") or group
    # Uniqueness within the *target* group only
    if (data["code"] != code or target_group != group) and \
            await db.accounts.select_one({"code": data["code"], "group": target_group}):
        raise HTTPException(status_code=400, detail=f"Kode akun {data['code']} sudah ada di kelompok {target_group}")
    await db.accounts.modify_one({"code": code, "group": group}, {"set": data})
    updated = await db.accounts.select_one({"code": data["code"], "group": target_group}, {"_id": 0})
    return Account(**updated)

@router.delete("/accounts/{code}")
async def delete_account(code: str, group: str = "BUMDES",
                         _: dict = Depends(require_roles(*ADMIN_LEVEL))):
    # Only block delete if a tx *in this group* uses this code
    unit_filter: dict
    if group == "BUMDES":
        unit_filter = {"unit_usaha_id": None}
    else:
        unit_doc = await db.unit_usaha.select_one({"code": group}, {"_id": 0, "id": 1})
        unit_filter = {"unit_usaha_id": (unit_doc or {}).get("id")} if unit_doc else {"unit_usaha_id": "__none__"}
    used = await db.transactions.select_one(
        {"$and": [
            {"$or": [{"debit_account_code": code}, {"credit_account_code": code}]},
            unit_filter,
        ]},
        {"_id": 1},
    )
    if used:
        raise HTTPException(status_code=400, detail="Kode akun sedang dipakai transaksi, tidak dapat dihapus")
    r = await db.accounts.remove_one({"code": code, "group": group})
    if r.deleted_count == 0:
        raise HTTPException(status_code=404, detail=f"Kode akun {code} tidak ditemukan di kelompok {group}")
    return {"deleted": r.deleted_count}
