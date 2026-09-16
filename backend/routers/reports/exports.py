"""Modular routes: exports"""
from router_dependencies import (
    ACCESS_TOKEN_EXPIRE_HOURS,
    ADMIN_LEVEL,
    APIRouter,
    API_PREFIX,
    APP_TITLE,
    Account,
    AccountCreate,
    DatabaseClient,
    CHART_OF_ACCOUNTS,
    COOKIE_NAME,
    CORSMiddleware,
    ChangePasswordRequest,
    Depends,
    FastAPI,
    File,
    HTMLResponse,
    HTTPException,
    List,
    Mitra,
    MitraCreate,
    Optional,
    P,
    Paragraph,
    PasswordResetRequest,
    Path,
    ProfileUpdateRequest,
    Query,
    READONLY_ROLES,
    READ_LEVEL,
    REPORT_READ_LEVEL,
    ROOT_DIR,
    Request,
    Response,
    RevenueShare,
    RevenueShareCreate,
    Spacer,
    StreamingResponse,
    Table,
    TableStyle,
    Transaction,
    TransactionCreate,
    UNIT_USAHA_SEED,
    UnitUsaha,
    UnitUsahaCreate,
    UploadFile,
    User,
    UserCreate,
    UserLogin,
    UserOut,
    UserRole,
    WRITE_LEVEL,
    _CELL_BOLD,
    _CELL_BOLD_RIGHT,
    _CELL_RIGHT,
    _CELL_STYLE,
    _STYLES,
    _arus_kas,
    _calc_balances,
    _calc_balances_before,
    _get_accounts_map,
    _group_from_unit,
    _laba_rugi,
    _ledger_data,
    _neraca,
    _pdf_response,
    _per_unit_report,
    _perubahan_ekuitas,
    _section_row,
    _sig_flow,
    _table_style,
    app,
    client,
    close_database,
    cm,
    colors,
    create_access_token,
    datetime,
    db,
    fmt_rp,
    get_current_user_payload,
    hash_password,
    io,
    load_dotenv,
    logging,
    now_utc,
    os,
    pdf_header,
    require_not_readonly,
    require_password_ready,
    require_roles,
    router,
    scope_unit_for_pengelola,
    seed_database,
    timezone,
    user_from_payload,
    uuid,
    verify_password
)

router = APIRouter(prefix=API_PREFIX)

@router.get("/reports/{report_type}/excel")
async def export_report_excel(
    report_type: str,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    as_of_date: Optional[str] = None,
    unit_usaha_id: Optional[str] = None,
    account_code: Optional[str] = None,
    payload: dict = Depends(require_roles(*REPORT_READ_LEVEL)),
):
    unit_usaha_id = await scope_unit_for_pengelola(payload, unit_usaha_id)
    """Universal Excel export untuk semua jenis laporan."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    header_fill = PatternFill(start_color="D4E09B", end_color="D4E09B", fill_type="solid")
    section_fill = PatternFill(start_color="F5F1E8", end_color="F5F1E8", fill_type="solid")
    total_fill = PatternFill(start_color="D4E09B", end_color="D4E09B", fill_type="solid")
    bold = Font(bold=True)
    right = Alignment(horizontal="right")

    def hdr(title, subtitle=""):
        ws.append(["BUMDES KARYA RAHARJA"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["Desa Wonoharjo, Kecamatan Pangandaran, Kabupaten Pangandaran"])
        ws.append([title.upper()])
        ws["A3"].font = Font(bold=True, size=12)
        if subtitle:
            ws.append([subtitle])
        ws.append([])

    def add_signature():
        pass  # signature ditulis inline di akhir builder

    if report_type == "laba-rugi":
        data = await _laba_rugi(start_date, end_date, unit_usaha_id)
        hdr("Laporan Laba Rugi", f"Periode: {start_date} s.d. {end_date}")
        ws.append(["Kode", "Nama Akun", "Jumlah (Rp)"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        ws.append(["", "PENDAPATAN", ""]); ws[ws.max_row][1].font = bold
        for it in data["pendapatan"]:
            ws.append([it["code"], it["name"], it["amount"]])
        ws.append(["", "Total Pendapatan", data["total_pendapatan"]]); ws[ws.max_row][1].font = bold
        ws.append(["", "BEBAN", ""]); ws[ws.max_row][1].font = bold
        for it in data["beban"]:
            ws.append([it["code"], it["name"], it["amount"]])
        ws.append(["", "Total Beban", data["total_beban"]]); ws[ws.max_row][1].font = bold
        ws.append(["", "LABA / (RUGI) BERSIH", data["laba_bersih"]])
        for c in ws[ws.max_row]: c.font = bold; c.fill = total_fill
    elif report_type == "neraca":
        data = await _neraca(as_of_date, unit_usaha_id)
        hdr("Neraca", f"Per: {as_of_date}")
        ws.append(["Kode", "Akun", "Jumlah (Rp)"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        for section, items, total, label in [
            ("ASET", data["aset"], data["total_aset"], "Total Aset"),
            ("KEWAJIBAN", data["kewajiban"], data["total_kewajiban"], "Total Kewajiban"),
            ("EKUITAS", data["ekuitas"], data["total_ekuitas"], "Total Ekuitas"),
        ]:
            ws.append(["", section, ""]); ws[ws.max_row][1].font = bold
            for it in items:
                ws.append([it["code"], it["name"], it["amount"]])
            ws.append(["", label, total]); ws[ws.max_row][1].font = bold
        ws.append(["", "TOTAL PASIVA (Kewajiban + Ekuitas)", data["total_pasiva"]])
        for c in ws[ws.max_row]: c.font = bold; c.fill = total_fill
    elif report_type == "arus-kas":
        data = await _arus_kas(start_date, end_date, unit_usaha_id)
        hdr("Laporan Arus Kas", f"Periode: {start_date} s.d. {end_date}")
        ws.append(["Tanggal", "Keterangan", "Jumlah (Rp)"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        ws.append(["", "KAS MASUK", ""]); ws[ws.max_row][1].font = bold
        for it in data["kas_masuk"]:
            ws.append([it["date"], it["description"], it["amount"]])
        ws.append(["", "Total Kas Masuk", data["total_masuk"]]); ws[ws.max_row][1].font = bold
        ws.append(["", "KAS KELUAR", ""]); ws[ws.max_row][1].font = bold
        for it in data["kas_keluar"]:
            ws.append([it["date"], it["description"], it["amount"]])
        ws.append(["", "Total Kas Keluar", data["total_keluar"]]); ws[ws.max_row][1].font = bold
        ws.append(["", "ARUS KAS BERSIH", data["arus_kas_bersih"]])
        for c in ws[ws.max_row]: c.font = bold; c.fill = total_fill
    elif report_type == "perubahan-ekuitas":
        data = await _perubahan_ekuitas(start_date, end_date, unit_usaha_id)
        hdr("Laporan Perubahan Ekuitas", f"Periode: {start_date} s.d. {end_date}")
        ws.append(["No.", "Uraian", "Jumlah (Rp)"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        for item in data.get("rows", []):
            ws.append([item["no"], item["label"], "" if item.get("amount") is None else item["amount"]])
            if item.get("kind") == "section":
                for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
            elif item.get("bold"):
                for c in ws[ws.max_row]: c.font = bold; c.fill = total_fill
    elif report_type == "calk":
        if unit_usaha_id:
            raise HTTPException(status_code=404, detail="CALK hanya tersedia untuk grup BUMDES")
        laba = await _laba_rugi(start_date, end_date)
        neraca = await _neraca(end_date)
        arus_kas = await _arus_kas(start_date, end_date)
        hdr("Catatan atas Laporan Keuangan (CaLK)", f"Periode: {start_date} s.d. {end_date}")
        ws.append(["Bagian", "Uraian", "Jumlah (Rp)"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        ws.append(["1", "Informasi umum", ""])
        ws.append(["", "Nama", "BUMDES Karya Raharja"])
        ws.append(["", "Alamat", "Desa Wonoharjo, Kecamatan Pangandaran"])
        ws.append(["", "Dasar hukum", "Kepmendesa PDTT No. 136 Tahun 2022"])
        ws.append(["2", "Ringkasan kinerja", ""])
        for label, value in [
            ("Total pendapatan", laba["total_pendapatan"]),
            ("Total beban", laba["total_beban"]),
            ("Laba bersih", laba["laba_bersih"]),
            ("Total aset", neraca["total_aset"]),
            ("Total kewajiban", neraca["total_kewajiban"]),
            ("Total ekuitas", neraca["total_ekuitas"]),
            ("Arus kas bersih", arus_kas["arus_kas_bersih"]),
        ]:
            ws.append(["", label, value])
        ws.append(["3", "Kebijakan akuntansi", ""])
        for policy in [
            "Laporan disusun sesuai Kepmendesa PDTT No. 136 Tahun 2022.",
            "Pengakuan pendapatan menggunakan basis akrual.",
            "Bagi hasil pengelola sebesar 30% dari laba bersih unit usaha.",
            "Bagi hasil BUMDES sebesar 70% dari laba bersih unit usaha.",
        ]:
            ws.append(["", policy, ""])
    elif report_type == "per-unit":
        data = await _per_unit_report(start_date, end_date)
        b = data.get("bumdes") or {}
        hdr("Rekap Kinerja BUMDES & Unit Usaha", f"Periode: {start_date} s.d. {end_date}")
        # ==== Tabel Kinerja BUMDES ====
        ws.append(["Kinerja BUMDES"])
        for c in ws[ws.max_row]: c.font = bold
        ws.append(["Kode", "Pendapatan", "Beban", "Laba Bersih", "18% Modal", "82% Unsur Lain"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        alokasi_82 = (
            f"PADes (30%) = Rp {int(b.get('share_pades_30', 0)):,}\n"
            f"Penasihat (7%) = Rp {int(b.get('share_penasihat_7', 0)):,}\n"
            f"Pengawas (5%) = Rp {int(b.get('share_pengawas_5', 0)):,}\n"
            f"Pengurus (35%) = Rp {int(b.get('share_pengurus_35', 0)):,}\n"
            f"Dana Sosial (5%) = Rp {int(b.get('share_dana_sosial_5', 0)):,}"
        ).replace(",", ".")
        ws.append(["BUMDES", b.get("pendapatan", 0), b.get("beban", 0), b.get("laba_bersih", 0),
                   b.get("share_modal_18", 0), alokasi_82])
        from openpyxl.styles import Alignment
        ws.cell(row=ws.max_row, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 90
        ws.append([])
        # ==== Tabel Kinerja Per Unit ====
        ws.append(["Kinerja Per Unit Usaha"])
        for c in ws[ws.max_row]: c.font = bold
        ws.append(["Kode", "Unit Usaha", "Pendapatan", "Beban", "Laba Bersih", "30% Pengelola", "70% BUMDES"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        tot_p = tot_b = tot_l = tot_m = tot_bm = 0.0
        for u in data["units"]:
            ws.append([u["code"], u["name"], u["pendapatan"], u["beban"], u["laba_bersih"],
                       u["share_pengelola_30"], u["share_bumdes_70"]])
            tot_p += u["pendapatan"]; tot_b += u["beban"]; tot_l += u["laba_bersih"]
            tot_m += u["share_pengelola_30"]; tot_bm += u["share_bumdes_70"]
        ws.append(["", "TOTAL", tot_p, tot_b, tot_l, tot_m, tot_bm])
        for c in ws[ws.max_row]: c.font = bold; c.fill = total_fill
    elif report_type == "ledger":
        if not account_code:
            raise HTTPException(status_code=400, detail="account_code wajib")
        data = await _ledger_data(account_code, start_date, end_date, unit_usaha_id)
        hdr("Buku Besar", f"Akun: {data['account']['code']} - {data['account']['name']}")
        ws.append(["Tanggal", "Keterangan", "Akun Lawan", "Ref.", "Debit", "Kredit", "Saldo"])
        for c in ws[ws.max_row]: c.font = bold; c.fill = header_fill
        ws.append(["", "Saldo Awal", "", "", "", "", data["saldo_awal"]])
        for e in data["entries"]:
            ws.append([e["date"], e["description"] or "", f"{e['other_account_code']} - {e.get('other_account_name','')}",
                       e.get("reference") or "", e["debit"] or 0, e["credit"] or 0, e["balance"]])
        ws.append(["", "TOTAL PERIODE", "", "", data["total_debit"], data["total_credit"], data["saldo_akhir"]])
        for c in ws[ws.max_row]: c.font = bold; c.fill = total_fill
    else:
        raise HTTPException(status_code=400, detail="report_type tidak dikenal")

    # Column widths
    for col_idx in range(1, ws.max_column + 1):
        max_len = 10
        for r in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for c in r:
                if c.value is not None:
                    max_len = max(max_len, min(50, len(str(c.value)) + 2))
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len

    # Signature
    ws.append([])
    from datetime import datetime as _dt
    ws.append([f"Pangandaran, {_dt.now().strftime('%d %B %Y')}"])
    ws.append(["Mengetahui,", "", "", "", "Disusun oleh,"])
    ws.append(["Direktur BUMDES", "", "", "", "Bendahara BUMDES"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{report_type}.xlsx"'},
    )
