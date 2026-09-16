"""Modular routes: imports"""
from router_dependencies import (
    ACCESS_TOKEN_EXPIRE_HOURS,
    ADMIN_LEVEL,
    APIRouter,
    API_PREFIX,
    APP_TITLE,
    Account,
    AccountCreate,
    DatabaseClient,
    CHART_OF_ACCOUNTS,
    COOKIE_NAME,
    CORSMiddleware,
    ChangePasswordRequest,
    Depends,
    FastAPI,
    File,
    HTMLResponse,
    HTTPException,
    List,
    Mitra,
    MitraCreate,
    Optional,
    P,
    Paragraph,
    PasswordResetRequest,
    Path,
    ProfileUpdateRequest,
    Query,
    READONLY_ROLES,
    READ_LEVEL,
    REPORT_READ_LEVEL,
    ROOT_DIR,
    Request,
    Response,
    RevenueShare,
    RevenueShareCreate,
    Spacer,
    StreamingResponse,
    Table,
    TableStyle,
    Transaction,
    TransactionCreate,
    UNIT_USAHA_SEED,
    UnitUsaha,
    UnitUsahaCreate,
    UploadFile,
    User,
    UserCreate,
    UserLogin,
    UserOut,
    UserRole,
    WRITE_LEVEL,
    _CELL_BOLD,
    _CELL_BOLD_RIGHT,
    _CELL_RIGHT,
    _CELL_STYLE,
    _STYLES,
    _arus_kas,
    _calc_balances,
    _calc_balances_before,
    _get_accounts_map,
    _group_from_unit,
    _laba_rugi,
    _ledger_data,
    _neraca,
    _pdf_response,
    _per_unit_report,
    _perubahan_ekuitas,
    _section_row,
    _sig_flow,
    _table_style,
    app,
    client,
    close_database,
    cm,
    colors,
    create_access_token,
    datetime,
    db,
    fmt_rp,
    get_current_user_payload,
    hash_password,
    io,
    load_dotenv,
    logging,
    now_utc,
    os,
    pdf_header,
    require_not_readonly,
    require_password_ready,
    require_roles,
    router,
    scope_unit_for_pengelola,
    seed_database,
    timezone,
    user_from_payload,
    uuid,
    verify_password
)

router = APIRouter(prefix=API_PREFIX)

@router.get("/transactions/template")
async def download_transaction_template(_: dict = Depends(require_roles(*WRITE_LEVEL))):
    """Generate template Excel untuk impor transaksi."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaksi"
    headers = ["tanggal", "unit_code", "jenis_transaksi", "keterangan", "nominal", "debit", "kredit", "referensi"]
    ws.append(headers)
    # Sample rows
    ws.append(["2025-06-01", "UU01", "penerimaan_bagi_hasil_domba", "Contoh: bagi hasil domba", 500000, "", "", "nota-001"])
    ws.append(["2025-06-05", "UU02", "penerimaan_setoran_ikan", "Contoh: setoran Rp3.000/kg", 150000, "", "", ""])
    ws.append(["2025-06-10", "", "modal_masuk", "Contoh: penyertaan modal desa", 5000000, "", "", ""])
    # Column widths
    widths = [12, 10, 32, 32, 14, 10, 10, 14]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # Second sheet: reference codes
    ref = wb.create_sheet("Referensi")
    ref.append(["Kode Unit", "Nama Unit"])
    for u in await db.unit_usaha.select({}, {"_id": 0}).order("code", 1).all(50):
        ref.append([u["code"], u["name"]])
    ref.append([])
    ref.append(["Kode Jenis Transaksi", "Nama", "Berlaku untuk Unit"])
    for t in await db.transaction_types.select({}, {"_id": 0}).all(500):
        ref.append([t["code"], t["name"], ", ".join(t.get("unit_codes", [])) or "Umum"])
    ref.column_dimensions["A"].width = 40
    ref.column_dimensions["B"].width = 50
    ref.column_dimensions["C"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-Transaksi-BUMDES.xlsx"'},
    )

@router.post("/transactions/import")
async def import_transactions_excel(
    file: UploadFile = File(...),
    dep: dict = Depends(require_roles(*WRITE_LEVEL)),
):
    """Import transactions from Excel. Expected columns:
    tanggal (YYYY-MM-DD), unit_code (opsional), jenis_transaksi (code),
    keterangan, nominal, debit (opsional), kredit (opsional), referensi (opsional).
    """
    from decimal import Decimal, InvalidOperation
    from openpyxl import load_workbook
    user = await user_from_payload(dep)
    contents = await file.read()
    try:
        wb = load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"File Excel tidak valid: {e}")
    ws = wb.active

    header = [str(c.value or "").strip().lower() for c in ws[1]]

    def col(name):
        try:
            return header.index(name)
        except ValueError:
            return None

    idx_tanggal = col("tanggal")
    idx_unit = col("unit_code")
    idx_type = col("jenis_transaksi")
    idx_desc = col("keterangan")
    idx_amt = col("nominal")
    idx_deb = col("debit")
    idx_kre = col("kredit")
    idx_ref = col("referensi")

    if idx_tanggal is None or idx_type is None or idx_amt is None:
        raise HTTPException(
            status_code=400,
            detail="Kolom minimal wajib: tanggal, jenis_transaksi, nominal (lihat template)."
        )

    # Preload
    tx_types = {t["code"]: t for t in await db.transaction_types.select({}, {"_id": 0}).all(500)}
    units_by_code = {u["code"]: u for u in await db.unit_usaha.select({}, {"_id": 0}).all(50)}

    inserted_docs, errors = [], []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v in (None, "") for v in row):
            continue
        try:
            tanggal = row[idx_tanggal]
            if hasattr(tanggal, "strftime"):
                tanggal = tanggal.strftime("%Y-%m-%d")
            tanggal = str(tanggal).strip()

            type_code = str(row[idx_type]).strip() if row[idx_type] else ""
            if not type_code or type_code not in tx_types:
                raise ValueError(f"Jenis transaksi '{type_code}' tidak dikenal")
            tt = tx_types[type_code]

            try:
                amount = Decimal(str(row[idx_amt] or 0)).quantize(Decimal("0.01"))
            except (InvalidOperation, ValueError) as exc:
                raise ValueError(f"Nominal tidak valid: {exc}")
            desc = str(row[idx_desc]).strip() if idx_desc is not None and row[idx_desc] else tt["name"]
            unit_code = str(row[idx_unit]).strip() if idx_unit is not None and row[idx_unit] else ""
            unit_id = units_by_code[unit_code]["id"] if unit_code and unit_code in units_by_code else None
            debit = str(row[idx_deb]).strip() if idx_deb is not None and row[idx_deb] else tt["debit"]
            kredit = str(row[idx_kre]).strip() if idx_kre is not None and row[idx_kre] else tt["credit"]
            ref = str(row[idx_ref]).strip() if idx_ref is not None and row[idx_ref] else ""

            tx = Transaction(
                date=tanggal, unit_usaha_id=unit_id, transaction_type=type_code,
                description=desc, amount=amount,
                debit_account_code=debit, credit_account_code=kredit,
                reference=ref, created_by=user.id,
            )
            inserted_docs.append(tx.model_dump())
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})

    if inserted_docs:
        await db.transactions.create_many(inserted_docs)
    return {"inserted": len(inserted_docs), "errors": errors, "total_rows": ws.max_row - 1}
