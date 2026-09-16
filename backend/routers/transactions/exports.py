"""Modular routes: exports"""
from router_dependencies import (
    ACCESS_TOKEN_EXPIRE_HOURS,
    ADMIN_LEVEL,
    APIRouter,
    API_PREFIX,
    APP_TITLE,
    Account,
    AccountCreate,
    DatabaseClient,
    CHART_OF_ACCOUNTS,
    COOKIE_NAME,
    CORSMiddleware,
    ChangePasswordRequest,
    Depends,
    FastAPI,
    File,
    HTMLResponse,
    HTTPException,
    List,
    Mitra,
    MitraCreate,
    Optional,
    P,
    Paragraph,
    PasswordResetRequest,
    Path,
    ProfileUpdateRequest,
    Query,
    READONLY_ROLES,
    READ_LEVEL,
    REPORT_READ_LEVEL,
    ROOT_DIR,
    Request,
    Response,
    RevenueShare,
    RevenueShareCreate,
    Spacer,
    StreamingResponse,
    Table,
    TableStyle,
    Transaction,
    TransactionCreate,
    UNIT_USAHA_SEED,
    UnitUsaha,
    UnitUsahaCreate,
    UploadFile,
    User,
    UserCreate,
    UserLogin,
    UserOut,
    UserRole,
    WRITE_LEVEL,
    _CELL_BOLD,
    _CELL_BOLD_RIGHT,
    _CELL_RIGHT,
    _CELL_STYLE,
    _STYLES,
    _arus_kas,
    _calc_balances,
    _calc_balances_before,
    _get_accounts_map,
    _group_from_unit,
    _laba_rugi,
    _ledger_data,
    _neraca,
    _pdf_response,
    _per_unit_report,
    _perubahan_ekuitas,
    _section_row,
    _sig_flow,
    _table_style,
    app,
    client,
    close_database,
    cm,
    colors,
    create_access_token,
    datetime,
    db,
    fmt_rp,
    get_current_user_payload,
    hash_password,
    io,
    load_dotenv,
    logging,
    now_utc,
    os,
    pdf_header,
    require_not_readonly,
    require_password_ready,
    require_roles,
    router,
    scope_unit_for_pengelola,
    seed_database,
    timezone,
    user_from_payload,
    uuid,
    verify_password
)

router = APIRouter(prefix=API_PREFIX)

@router.get("/transactions/export")
async def export_transactions_excel(
    unit_usaha_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    all_data: bool = False,
    payload: dict = Depends(get_current_user_payload),
):
    """Export transaksi ke Excel.
    - Default: filter by unit_usaha_id + tanggal (single-sheet).
    - all_data=true: multi-sheet workbook berisi SEMUA transaksi (BUMDES + tiap unit + sheet gabungan).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    title_font = Font(bold=True, size=13, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="D4E09B")
    title_fill = PatternFill("solid", fgColor="2E4F32")
    total_fill = PatternFill("solid", fgColor="F5F1E8")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="C8CFC1")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    accounts = {a["code"]: a for a in await db.accounts.select({}, {"_id": 0}).all(500)}
    units = {u["id"]: u for u in await db.unit_usaha.select({}, {"_id": 0}).all(50)}
    widths = [12, 10, 28, 40, 12, 32, 12, 32, 14, 14]
    headers = ["Tanggal", "Kelompok", "Jenis Transaksi", "Keterangan",
               "Kode Debit", "Nama Debit", "Kode Kredit", "Nama Kredit",
               "Nominal", "Referensi"]

    def _amount(value):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    def _row_from_tx(t: dict):
        uid = t.get("unit_usaha_id")
        grp = units.get(uid, {}).get("code") if uid else "BUMDES"
        d_code = t.get("debit_account_code", ""); c_code = t.get("credit_account_code", "")
        return [
            t.get("date", ""), grp,
            t.get("transaction_type", "") or "",
            t.get("description", "") or "",
            d_code, accounts.get(d_code, {}).get("name", ""),
            c_code, accounts.get(c_code, {}).get("name", ""),
            _amount(t.get("amount")),
            t.get("reference", "") or "",
        ]

    def _fill_sheet(ws, title_text: str, txs: list):
        ws.append([title_text])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
        tc = ws.cell(row=1, column=1)
        tc.font = title_font; tc.fill = title_fill; tc.alignment = center
        ws.row_dimensions[1].height = 28
        ws.append([f"Jumlah transaksi: {len(txs)}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=10)
        ws.cell(row=2, column=1).font = Font(italic=True, color="5C6E5E")
        ws.append(headers)
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = header_fill; c.alignment = center; c.border = border_all
        if not txs:
            ws.append(["", "", "", "Tidak ada transaksi.", "", "", "", "", "", ""])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=10)
            ws.cell(row=ws.max_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="8A9A8C")
        else:
            total = 0
            for t in txs:
                ws.append(_row_from_tx(t))
                total += _amount(t.get("amount"))
            ws.append(["", "", "", "TOTAL", "", "", "", "", total, ""])
            for c in ws[ws.max_row]:
                c.font = bold; c.fill = total_fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"

    wb = Workbook()

    if all_data:
        # Pengelola hanya boleh export data unitnya sendiri
        if payload.get("role") == "pengelola":
            uid = payload.get("unit") or None
            all_txs = await db.transactions.select({"unit_usaha_id": uid}, {"_id": 0}).order("date", 1).all(50000)
            wb.remove(wb.active)
            unit_code = (units.get(uid) or {}).get("code") or "Unit"
            ws = wb.create_sheet(unit_code)
            _fill_sheet(ws, f"Transaksi {unit_code} — Semua Periode", all_txs)
            fname = f"Transaksi_{unit_code}_Semua_Data.xlsx"
        else:
            all_txs = await db.transactions.select({}, {"_id": 0}).order("date", 1).all(50000)
            # Sheet 1: Semua (gabungan sortir tanggal)
            ws0 = wb.active; ws0.title = "Semua"
            _fill_sheet(ws0, "Semua Transaksi — BUMDES & Unit Usaha", all_txs)
            # Sheet BUMDES
            ws_b = wb.create_sheet("BUMDES")
            _fill_sheet(ws_b, "Transaksi BUMDES — Semua Periode",
                        [t for t in all_txs if not t.get("unit_usaha_id")])
            # Sheet per unit (urutan berdasar code)
            for u in sorted(units.values(), key=lambda x: x.get("code") or ""):
                ws_u = wb.create_sheet(u["code"])
                _fill_sheet(ws_u, f"Transaksi {u['code']} · {u['name']} — Semua Periode",
                            [t for t in all_txs if t.get("unit_usaha_id") == u["id"]])
            fname = f"Transaksi_Semua_Data_{datetime.now().strftime('%Y%m%d')}.xlsx"
    else:
        # Single-sheet: existing behaviour with unit_usaha_id + date range filter
        unit_usaha_id = await scope_unit_for_pengelola(payload, unit_usaha_id)
        q: dict = {}
        if unit_usaha_id == "":
            q["unit_usaha_id"] = None
        elif unit_usaha_id is not None:
            q["unit_usaha_id"] = unit_usaha_id
        if start_date and end_date:
            q["date"] = {"$gte": start_date, "$lte": end_date}
        txs = await db.transactions.select(q, {"_id": 0}).order("date", 1).all(20000)
        scope_lbl = "BUMDES" if unit_usaha_id == "" else (units.get(unit_usaha_id, {}).get("code") or "Semua")
        period_lbl = f"{start_date or '-'} s.d. {end_date or '-'}"
        ws = wb.active; ws.title = "Transaksi"
        _fill_sheet(ws, f"Transaksi Keuangan · Kelompok: {scope_lbl} · Periode: {period_lbl}", txs)
        fname = f"Transaksi_{scope_lbl}_{start_date or 'all'}_sd_{end_date or 'all'}.xlsx"

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
